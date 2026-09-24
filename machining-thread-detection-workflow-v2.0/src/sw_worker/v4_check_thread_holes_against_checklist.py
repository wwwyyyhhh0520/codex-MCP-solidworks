import argparse
import csv
import json
import os
import re
import tempfile
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


FIELD_ALIASES = {
    "part_key": ["part_key", "part", "零件", "零件号", "零件名称", "零件/实例", "孔零件/实例", "组件", "occurrence"],
    "feature_key": ["feature_key", "feature", "孔特征", "特征", "孔名称", "feature_name", "螺纹孔", "孔位"],
    "hole_spec": ["hole_spec", "spec", "规格", "孔规格", "螺纹规格", "thread_spec"],
    "expected_count": ["expected_count", "count", "数量", "要求数量", "应有数量", "标准数量", "qty"],
}


ISSUE_TEXT = {
    "CHECKLIST_ITEM_MISSING_IN_MODEL": "清单要求存在该螺纹孔，但模型审核结果中未找到对应记录",
    "CHECKLIST_COUNT_NOT_ENOUGH": "模型中对应螺纹孔数量少于清单要求",
    "CHECKLIST_SPEC_MISMATCH": "模型中存在同名孔特征，但螺纹规格与清单不一致",
    "MODEL_EXTRA_THREAD_HOLE": "模型中存在清单未列出的螺纹孔",
    "MODEL_ITEM_HAS_ACTIONABLE_RESULT": "清单项在模型中存在，但模型审核结论需要处理或复核",
}


ACTION_TEXT = {
    "CHECKLIST_ITEM_MISSING_IN_MODEL": "在模型中补建对应 HoleWizard 螺纹孔，或确认清单要求已取消",
    "CHECKLIST_COUNT_NOT_ENOUGH": "复核该零件/孔特征的数量，补齐缺少的螺纹孔实例",
    "CHECKLIST_SPEC_MISMATCH": "复核模型孔规格与清单规格，修正 HoleWizard 规格或清单",
    "MODEL_EXTRA_THREAD_HOLE": "确认该螺纹孔是否为新增设计；若非新增，应从模型或清单中修正",
    "MODEL_ITEM_HAS_ACTIONABLE_RESULT": "按模型审核结论处理自动缺陷或完成人工复核闭环",
}


def norm(value) -> str:
    return str(value or "").strip()


def norm_header(value) -> str:
    return re.sub(r"\s+", "", norm(value).lower())


def clean_part(value: str) -> str:
    value = norm(value).split("/")[-1]
    return re.sub(r"-\d+$", "", value)


def canonical_spec(value: str) -> str:
    value = norm(value).upper().replace(" ", "")
    m = re.search(r"M\d+(?:\.\d+)?", value)
    if m:
        return m.group(0)
    return value


def find_column(headers: list[str], key: str) -> str | None:
    normalized = {norm_header(h): h for h in headers}
    for alias in FIELD_ALIASES[key]:
        hit = normalized.get(norm_header(alias))
        if hit:
            return hit
    return None


def read_csv_rows(path: Path) -> list[dict[str, object]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def read_xlsx_rows(path: Path) -> list[dict[str, object]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [norm(v) for v in rows[0]]
    output = []
    for row in rows[1:]:
        if not any(v is not None and norm(v) for v in row):
            continue
        output.append({headers[i]: row[i] if i < len(row) else "" for i in range(len(headers))})
    return output


def read_table(path: Path) -> list[dict[str, object]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_csv_rows(path)
    if suffix in {".xlsx", ".xlsm"}:
        return read_xlsx_rows(path)
    raise ValueError(f"UNSUPPORTED_TABLE_TYPE={path}")


def normalize_checklist_rows(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    if not rows:
        return []
    headers = list(rows[0].keys())
    part_col = find_column(headers, "part_key")
    feature_col = find_column(headers, "feature_key")
    spec_col = find_column(headers, "hole_spec")
    count_col = find_column(headers, "expected_count")
    missing = [name for name, col in [("part_key", part_col), ("feature_key", feature_col), ("hole_spec", spec_col)] if not col]
    if missing:
        raise ValueError(f"CHECKLIST_REQUIRED_COLUMNS_MISSING={','.join(missing)}")

    normalized = []
    for idx, row in enumerate(rows, start=2):
        part = clean_part(row.get(part_col, ""))
        feature = norm(row.get(feature_col, ""))
        spec = canonical_spec(row.get(spec_col, ""))
        if not part and not feature and not spec:
            continue
        try:
            expected_count = int(float(norm(row.get(count_col, "1") if count_col else "1") or "1"))
        except Exception:
            expected_count = 1
        expected_count = max(expected_count, 1)
        normalized.append(
            {
                "part_key": part,
                "feature_key": feature,
                "hole_spec": spec,
                "expected_count": expected_count,
                "source_row": idx,
            }
        )
    return normalized


def is_tapped(row: dict[str, object]) -> bool:
    kind = norm(row.get("hole_kind")).lower()
    feature = norm(row.get("feature_name"))
    spec = canonical_spec(row.get("hole_spec"))
    return kind == "tapped" or "螺纹" in feature or bool(re.match(r"^M\d+", spec))


def result_status(row: dict[str, object]) -> tuple[str, str, str]:
    status = norm(row.get("corrected_status") or row.get("status"))
    code = norm(row.get("corrected_finding_code") or row.get("finding_code"))
    finding = norm(row.get("corrected_finding") or row.get("finding"))
    return status, code, finding


def model_key(row: dict[str, object], include_spec: bool = True) -> tuple[str, str, str] | tuple[str, str]:
    part = clean_part(row.get("occurrence", ""))
    feature = norm(row.get("feature_name"))
    spec = canonical_spec(row.get("hole_spec"))
    return (part, feature, spec) if include_spec else (part, feature)


def checklist_key(row: dict[str, object], include_spec: bool = True) -> tuple[str, str, str] | tuple[str, str]:
    part = clean_part(row.get("part_key", ""))
    feature = norm(row.get("feature_key"))
    spec = canonical_spec(row.get("hole_spec"))
    return (part, feature, spec) if include_spec else (part, feature)


def build_diffs(checklist_path: Path, model_csv: Path, include_extra: bool = True) -> tuple[list[dict[str, object]], dict[str, int]]:
    checklist = normalize_checklist_rows(read_table(checklist_path))
    model_rows = [r for r in read_csv_rows(model_csv) if is_tapped(r)]

    expected_map: dict[tuple[str, str, str], int] = defaultdict(int)
    expected_source_rows: dict[tuple[str, str, str], list[int]] = defaultdict(list)
    for row in checklist:
        key = checklist_key(row)
        expected_map[key] += int(row["expected_count"])
        expected_source_rows[key].append(int(row["source_row"]))

    model_map: dict[tuple[str, str, str], list[dict[str, object]]] = defaultdict(list)
    model_by_part_feature: dict[tuple[str, str], list[dict[str, object]]] = defaultdict(list)
    for row in model_rows:
        model_map[model_key(row)].append(row)
        model_by_part_feature[model_key(row, include_spec=False)].append(row)

    diffs: list[dict[str, object]] = []
    for key in sorted(expected_map):
        part, feature, spec = key
        expected_count = expected_map[key]
        actual_rows = model_map.get(key, [])
        actual_count = len(actual_rows)
        target_result = result_text(actual_rows[0]) if actual_rows else ""
        source_rows = ",".join(str(v) for v in expected_source_rows[key])

        if actual_count == 0:
            same_name = model_by_part_feature.get((part, feature), [])
            if same_name:
                actual_specs = ",".join(sorted({canonical_spec(r.get("hole_spec")) for r in same_name}))
                diff_type = "CHECKLIST_SPEC_MISMATCH"
                issue = f"清单要求 {spec}，模型同名孔实际规格为 {actual_specs}"
                actual_count = len(same_name)
                target_result = result_text(same_name[0])
            else:
                diff_type = "CHECKLIST_ITEM_MISSING_IN_MODEL"
                issue = ISSUE_TEXT[diff_type]
        elif actual_count < expected_count:
            diff_type = "CHECKLIST_COUNT_NOT_ENOUGH"
            issue = f"清单要求 {expected_count} 个，模型实际 {actual_count} 个，缺少 {expected_count - actual_count} 个"
        else:
            actionable = first_actionable(actual_rows)
            if actionable:
                diff_type = "MODEL_ITEM_HAS_ACTIONABLE_RESULT"
                status, code, finding = result_status(actionable)
                issue = f"清单项存在，但模型审核结论为 {status} / {code}：{finding}"
            else:
                continue

        diffs.append(
            {
                "part_key": part,
                "feature_key": feature,
                "hole_spec": spec,
                "checklist_count": expected_count,
                "model_count": actual_count,
                "issue": issue,
                "model_result": target_result,
                "action": ACTION_TEXT[diff_type],
                "source_rows": source_rows,
                "raw_diff_type": diff_type,
            }
        )

    if include_extra:
        for key in sorted(set(model_map) - set(expected_map)):
            rows = model_map[key]
            part, feature, spec = key
            diffs.append(
                {
                    "part_key": part,
                    "feature_key": feature,
                    "hole_spec": spec,
                    "checklist_count": 0,
                    "model_count": len(rows),
                    "issue": ISSUE_TEXT["MODEL_EXTRA_THREAD_HOLE"],
                    "model_result": result_text(rows[0]),
                    "action": ACTION_TEXT["MODEL_EXTRA_THREAD_HOLE"],
                    "source_rows": "",
                    "raw_diff_type": "MODEL_EXTRA_THREAD_HOLE",
                }
            )

    metrics = {
        "checklist_item_count": len(checklist),
        "checklist_key_count": len(expected_map),
        "model_tapped_row_count": len(model_rows),
    }
    return diffs, metrics


def result_text(row: dict[str, object]) -> str:
    status, code, _ = result_status(row)
    return f"{status} / {code}".strip(" /")


def first_actionable(rows: list[dict[str, object]]) -> dict[str, object] | None:
    for row in rows:
        status, code, _ = result_status(row)
        if status in {"ISSUE", "REVIEW_REQUIRED"} or code and code not in {"HOLE_SIDE_GEOMETRY_PASS", "THREAD_HOLE_OPEN_TO_CAVITY_PASS", "HOLE_SIDE_GEOMETRY_PASS_BY_MAIN_BORE_DEPTH"}:
            return row
    return None


def write_outputs(diffs: list[dict[str, object]], metrics: dict[str, int], output_root: Path) -> tuple[Path, Path]:
    out_dir = output_root / f"thread_checklist_compare_v4_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    out_dir.mkdir(parents=True, exist_ok=True)
    os.environ["TMP"] = str(out_dir)
    os.environ["TEMP"] = str(out_dir)
    tempfile.tempdir = str(out_dir)

    xlsx = out_dir / "thread_checklist_compare_v4.xlsx"
    json_path = out_dir / "thread_checklist_compare_v4.json"

    machine_headers = [
        "part_key",
        "feature_key",
        "hole_spec",
        "checklist_count",
        "model_count",
        "issue",
        "model_result",
        "action",
        "source_rows",
        "raw_diff_type",
    ]
    display_columns = [
        ("part_key", "零件"),
        ("feature_key", "孔特征"),
        ("hole_spec", "孔规格"),
        ("checklist_count", "清单数量"),
        ("model_count", "模型数量"),
        ("issue", "问题说明"),
        ("model_result", "模型审核结论"),
        ("action", "建议动作"),
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "清单比对"
    ws.append([title for _, title in display_columns])
    for diff in diffs:
        ws.append([diff.get(key, "") for key, _ in display_columns])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    missing_fill = PatternFill("solid", fgColor="FFF2CC")
    diff_fill = PatternFill("solid", fgColor="FCE4D6")
    extra_fill = PatternFill("solid", fgColor="E2F0D9")
    issue_fill = PatternFill("solid", fgColor="F4CCCC")
    thin = Side(style="thin", color="9E9E9E")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws.iter_rows(min_row=2):
        diff_type = diffs[row[0].row - 2].get("raw_diff_type") if row[0].row - 2 < len(diffs) else ""
        fill = {
            "CHECKLIST_ITEM_MISSING_IN_MODEL": missing_fill,
            "CHECKLIST_COUNT_NOT_ENOUGH": missing_fill,
            "CHECKLIST_SPEC_MISMATCH": diff_fill,
            "MODEL_EXTRA_THREAD_HOLE": extra_fill,
            "MODEL_ITEM_HAS_ACTIONABLE_RESULT": issue_fill,
        }.get(diff_type, diff_fill)
        for idx, cell in enumerate(row, start=1):
            cell.fill = fill
            cell.border = border
            if idx in {3, 4, 5}:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    for idx, width in enumerate([22, 26, 12, 12, 12, 54, 36, 52], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    summary = wb.create_sheet("Summary")
    summary.append(["metric", "count"])
    summary.append(["diff_count", len(diffs)])
    for k, v in metrics.items():
        summary.append([k, v])
    for k, v in Counter(str(d["raw_diff_type"]) for d in diffs).items():
        summary.append([k, v])
    for cell in summary[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.border = border
    for row in summary.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center" if cell.column == 2 else "left", vertical="center")
    summary.column_dimensions["A"].width = 44
    summary.column_dimensions["B"].width = 16

    wb.save(xlsx)
    json_path.write_text(json.dumps({"metrics": metrics, "headers": machine_headers, "diffs": diffs}, ensure_ascii=False, indent=2), encoding="utf-8")
    return xlsx, json_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--checklist", required=True, help="结构化清单 .csv/.xlsx，至少包含零件、孔特征、孔规格，可选数量")
    parser.add_argument("--model-csv", required=True, help="模型审核输出：孔侧几何检测明细_V4.csv 或主孔深修正版 CSV")
    parser.add_argument("--output-root", required=True)
    parser.add_argument("--no-extra", action="store_true", help="不输出模型中清单未列出的额外螺纹孔")
    args = parser.parse_args()

    diffs, metrics = build_diffs(Path(args.checklist), Path(args.model_csv), include_extra=not args.no_extra)
    xlsx, json_path = write_outputs(diffs, metrics, Path(args.output_root))

    print(f"CHECKLIST_COMPARE_DIFF_COUNT={len(diffs)}")
    print("CHECKLIST_COMPARE_TYPE_COUNTS=" + json.dumps(dict(Counter(str(d["raw_diff_type"]) for d in diffs)), ensure_ascii=False))
    for key, value in metrics.items():
        print(f"{key.upper()}={value}")
    print(f"XLSX_PATH={xlsx}")
    print(f"JSON_PATH={json_path}")
    print("THREAD_CHECKLIST_COMPARE_V4_STATUS=SUCCESS")


if __name__ == "__main__":
    main()
