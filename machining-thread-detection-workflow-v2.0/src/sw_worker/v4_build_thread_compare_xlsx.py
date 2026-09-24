import argparse
import csv
import json
import os
import re
import tempfile
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ISSUE_TEXT = {
    "BASELINE_THREAD_HOLE_MISSING_IN_TARGET": "\u76ee\u6807\u65e0\u87ba\u7259\u6a21\u578b\u7f3a\u5c11\u8be5\u87ba\u7eb9\u5b54\u7279\u5f81",
    "TARGET_THREAD_HOLE_ONLY": "\u76ee\u6807\u6a21\u578b\u5b58\u5728\u8be5\u87ba\u7eb9\u5b54\uff0c\u57fa\u51c6\u6709\u87ba\u7259\u6a21\u578b\u65e0\u5bf9\u5e94\u8bb0\u5f55",
    "BOTH_EXIST_BUT_RESULT_DIFFERENT": "\u4e24\u7248\u6a21\u578b\u540c\u540d\u87ba\u7eb9\u5b54\u68c0\u6d4b\u7ed3\u8bba\u4e0d\u540c",
}

ACTION_TEXT = {
    "BASELINE_THREAD_HOLE_MISSING_IN_TARGET": "\u8bf7\u5728\u76ee\u6807\u6a21\u578b\u4e2d\u8865\u5efa\u5bf9\u5e94 HoleWizard \u87ba\u7eb9\u5b54\uff0c\u6216\u786e\u8ba4\u8be5\u87ba\u7eb9\u5b54\u5df2\u88ab\u8bbe\u8ba1\u53d6\u6d88",
    "TARGET_THREAD_HOLE_ONLY": "\u590d\u6838\u8be5\u5b54\u662f\u5426\u4e3a\u65b0\u589e\u8bbe\u8ba1\u6216\u57fa\u51c6\u6a21\u578b\u7f3a\u5931",
    "BOTH_EXIST_BUT_RESULT_DIFFERENT": "\u590d\u6838\u76ee\u6807\u6a21\u578b\u8be5\u5b54\u51e0\u4f55\u3001\u653b\u4e1d\u6df1\u5ea6\u3001\u5b54\u8f74\u548c\u5e95\u5b54\u5c3a\u5bf8",
}


def read_rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def clean_part(occurrence: str) -> str:
    value = str(occurrence or "").strip().split("/")[-1]
    return re.sub(r"-\d+$", "", value)


def is_tapped(row: dict[str, str]) -> bool:
    kind = str(row.get("hole_kind") or "").lower()
    feature = str(row.get("feature_name") or "")
    spec = str(row.get("hole_spec") or "").upper()
    return kind == "tapped" or "\u87ba\u7eb9" in feature or bool(re.match(r"^M\d+", spec))


def result_text(row: dict[str, str]) -> str:
    status = row.get("corrected_status") or row.get("status") or ""
    code = row.get("corrected_finding_code") or row.get("finding_code") or ""
    return f"{status} / {code}".strip(" /")


STATUS_ZH = {
    "PASS": "通过",
    "ISSUE": "存在问题",
    "REVIEW_REQUIRED": "需要复核",
    "NOT_APPLICABLE": "不适用",
}


FINDING_ZH = {
    "HOLE_SIDE_GEOMETRY_PASS": "孔侧几何通过",
    "HOLE_SIDE_GEOMETRY_PASS_BY_MAIN_BORE_DEPTH": "按主孔深修正后通过",
    "THREAD_HOLE_OPEN_TO_CAVITY_PASS": "螺纹孔开口到空腔，跳过盲孔底余量检查",
    "THREAD_DEPTH_EXCEEDS_MAIN_BORE_DEPTH": "攻丝深度超过主孔有效深度",
    "THREAD_DEPTH_EXCEEDS_PHYSICAL_HOLE_DEPTH": "攻丝深度超过检测物理孔深",
    "THREAD_DEPTH_EXCEEDS_MEASURED_DEPTH_BUT_BOTTOM_UNCONFIRMED": "攻丝深度超过可测孔深，但孔底未确认",
    "HOLE_AXIS_NOT_VERIFIED": "孔轴未可靠识别，需要复核",
    "TAP_DRILL_DIAMETER_DEVIATION_LARGE": "底孔直径偏差较大",
    "HOLE_AXIS_UNAVAILABLE": "孔轴不可读取",
    "HOLE_DIAMETER_NOT_MEASURABLE": "孔径不可测",
    "PHYSICAL_HOLE_DEPTH_NOT_MEASURABLE": "物理孔深不可测",
    "BLIND_HOLE_BOTTOM_CLEARANCE_SMALL": "盲孔底部余量不足",
}


def engineer_result_text(row: dict[str, str]) -> str:
    if not row:
        return "无对应记录"
    status = str(row.get("corrected_status") or row.get("status") or "").strip()
    code_text = str(row.get("corrected_finding_code") or row.get("finding_code") or "").strip()
    status_zh = STATUS_ZH.get(status, status or "未识别")
    codes = [c.strip() for c in re.split(r"[;,]", code_text) if c.strip()]
    if not codes:
        return status_zh
    translated = [FINDING_ZH.get(code, code) for code in codes]
    return status_zh + "：" + "；".join(translated)


def row_key(row: dict[str, str]) -> tuple[str, str, str]:
    return (
        clean_part(row.get("occurrence", "")),
        row.get("feature_name") or "",
        row.get("hole_spec") or "",
    )


def build_diffs(target_csv: Path, baseline_csv: Path) -> tuple[list[dict[str, object]], int, int]:
    target_rows = [r for r in read_rows(target_csv) if is_tapped(r)]
    baseline_rows = [r for r in read_rows(baseline_csv) if is_tapped(r)]

    target_map: dict[tuple[str, str, str], list[dict[str, str]]] = defaultdict(list)
    baseline_map: dict[tuple[str, str, str], list[dict[str, str]]] = defaultdict(list)
    for row in target_rows:
        target_map[row_key(row)].append(row)
    for row in baseline_rows:
        baseline_map[row_key(row)].append(row)

    diffs: list[dict[str, object]] = []
    for key in sorted(set(target_map) | set(baseline_map)):
        target = target_map.get(key, [])
        baseline = baseline_map.get(key, [])
        part, feature, spec = key
        target_result = result_text(target[0]) if target else ""
        baseline_result = result_text(baseline[0]) if baseline else ""
        target_result_zh = engineer_result_text(target[0]) if target else "无对应螺纹孔"
        baseline_result_zh = engineer_result_text(baseline[0]) if baseline else "无对应螺纹孔"

        target_status = str((target[0].get("corrected_status") or target[0].get("status") or "") if target else "").strip()
        baseline_status = str((baseline[0].get("corrected_status") or baseline[0].get("status") or "") if baseline else "").strip()

        if not target and baseline:
            diff_type = "BASELINE_THREAD_HOLE_MISSING_IN_TARGET"
        elif target and not baseline:
            diff_type = "TARGET_THREAD_HOLE_ONLY"
        elif target_status == "PASS" and baseline_status == "PASS":
            # 两边工程结论都是通过时，不进入工程师差异主表。
            # 例如“开口到空腔通过”和“孔侧几何通过”只是通过原因不同，不是需要处理的差异。
            continue
        elif target_result != baseline_result:
            diff_type = "BOTH_EXIST_BUT_RESULT_DIFFERENT"
        else:
            continue

        diffs.append(
            {
                "part_key": part,
                "feature_key": feature,
                "hole_spec": spec,
                "no_thread_count": len(target),
                "with_thread_count": len(baseline),
                "issue": ISSUE_TEXT[diff_type],
                "target_result": target_result_zh,
                "baseline_result": baseline_result_zh,
                "target_result_raw": target_result,
                "baseline_result_raw": baseline_result,
                "action": ACTION_TEXT[diff_type],
                "raw_diff_type": diff_type,
            }
        )

    return diffs, len(target_rows), len(baseline_rows)


def write_workbook(diffs: list[dict[str, object]], target_count: int, baseline_count: int, out_dir: Path) -> tuple[Path, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    os.environ["TMP"] = str(out_dir)
    os.environ["TEMP"] = str(out_dir)
    tempfile.tempdir = str(out_dir)
    xlsx = out_dir / "thread_hole_diff_v4.xlsx"
    json_path = out_dir / "thread_hole_diff_v4.json"

    machine_headers = [
        "part_key",
        "feature_key",
        "hole_spec",
        "no_thread_count",
        "with_thread_count",
        "issue",
        "target_result",
        "baseline_result",
        "action",
        "raw_diff_type",
    ]
    display_columns = [
        ("part_key", "零件"),
        ("feature_key", "孔特征"),
        ("hole_spec", "孔规格"),
        ("no_thread_count", "无螺牙模型中的数量"),
        ("with_thread_count", "有螺牙模型中的数量"),
        ("issue", "问题说明"),
        ("target_result", "无螺牙模型结论"),
        ("baseline_result", "有螺牙基准结论"),
        ("action", "建议动作"),
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "有无螺牙对比"
    ws.append([title for _, title in display_columns])

    for diff in diffs:
        ws.append([diff.get(key, "") for key, _ in display_columns])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    missing_fill = PatternFill("solid", fgColor="FFF2CC")
    diff_fill = PatternFill("solid", fgColor="E2F0D9")
    review_fill = PatternFill("solid", fgColor="FCE4D6")
    issue_fill = PatternFill("solid", fgColor="F4CCCC")
    thin = Side(style="thin", color="9E9E9E")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws.iter_rows(min_row=2):
        diff_type = diffs[row[0].row - 2].get("raw_diff_type") if row[0].row - 2 < len(diffs) else ""
        fill = missing_fill if diff_type == "BASELINE_THREAD_HOLE_MISSING_IN_TARGET" else diff_fill
        for idx, cell in enumerate(row, start=1):
            cell.fill = fill
            # 结论列两边共用同一套颜色规则：存在问题 > 需要复核 > 通过。
            if idx in {7, 8}:
                text_value = str(cell.value or "")
                if "存在问题" in text_value:
                    cell.fill = issue_fill
                elif "需要复核" in text_value:
                    cell.fill = review_fill
            cell.border = border
            if idx in {3, 4, 5}:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    for idx, width in enumerate([22, 26, 12, 16, 16, 42, 38, 38, 52], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    summary = wb.create_sheet("Summary")
    summary.append(["metric", "count"])
    summary.append(["diff_count", len(diffs)])
    for key, value in Counter(str(d["raw_diff_type"]) for d in diffs).items():
        summary.append([key, value])
    summary.append(["target_tapped_rows", target_count])
    summary.append(["baseline_tapped_rows", baseline_count])

    for cell in summary[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.border = border
    for row in summary.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center" if cell.column == 2 else "left", vertical="center")

    summary.column_dimensions["A"].width = 44
    summary.column_dimensions["B"].width = 16

    wb.save(xlsx)
    json_path.write_text(json.dumps({"headers": machine_headers, "diffs": diffs}, ensure_ascii=False, indent=2), encoding="utf-8")
    return xlsx, json_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--target-csv", required=True)
    parser.add_argument("--baseline-csv", required=True)
    parser.add_argument("--output-root", required=True)
    args = parser.parse_args()

    target_csv = Path(args.target_csv).resolve()
    baseline_csv = Path(args.baseline_csv).resolve()
    output_root = Path(args.output_root).resolve()
    out_dir = output_root / f"thread_compare_{datetime.now().strftime('%Y%m%d_%H%M%S')}"

    diffs, target_count, baseline_count = build_diffs(target_csv, baseline_csv)
    xlsx, json_path = write_workbook(diffs, target_count, baseline_count, out_dir)

    print(f"DIFF_COUNT={len(diffs)}")
    print("DIFF_TYPE_COUNTS=" + json.dumps(dict(Counter(str(d["raw_diff_type"]) for d in diffs)), ensure_ascii=False))
    print(f"TARGET_TAPPED_ROWS={target_count}")
    print(f"BASELINE_TAPPED_ROWS={baseline_count}")
    print(f"XLSX_PATH={xlsx}")
    print(f"JSON_PATH={json_path}")
    print("THREAD_COMPARE_XLSX_V4_STATUS=SUCCESS")


if __name__ == "__main__":
    main()
